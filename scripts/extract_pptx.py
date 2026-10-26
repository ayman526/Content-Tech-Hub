"""
extract_pptx.py — Extract translatable text segments from a PPTX file.

Outputs JSON to stdout:
{
  "segments": [
    {
      "slideIndex": 0,
      "shapeIndex": 0,
      "paragraphIndex": 0,
      "text": "Original text",
      "formatHint": "p"
    },
    ...
  ],
  "slideCount": N,
  "language": "en"
}

Usage:
  python3 scripts/extract_pptx.py <input.pptx>
"""

import sys
import json
import os

def extract_pptx(pptx_path: str) -> dict:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN

    prs = Presentation(pptx_path)
    segments = []
    slide_count = len(prs.slides)

    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if not shape.has_text_frame:
                continue
            for para_idx, para in enumerate(shape.text_frame.paragraphs):
                full_text = para.text.strip()
                if len(full_text) < 2:
                    continue

                # Determine format hint from paragraph level/font size
                fmt = "p"
                if para.runs:
                    run = para.runs[0]
                    size_pt = None
                    try:
                        size_pt = run.font.size.pt if run.font.size else None
                    except Exception:
                        pass
                    if size_pt is not None:
                        if size_pt >= 36:
                            fmt = "h1"
                        elif size_pt >= 28:
                            fmt = "h2"
                        elif size_pt >= 22:
                            fmt = "h3"
                        elif size_pt <= 12:
                            fmt = "p"

                # Check for list-like bullets
                if para.level > 0:
                    fmt = "li"

                segments.append({
                    "slideIndex": slide_idx,
                    "shapeIndex": shape_idx,
                    "paragraphIndex": para_idx,
                    "text": full_text,
                    "formatHint": fmt,
                })

    return {
        "segments": segments,
        "slideCount": slide_count,
        "language": "en",
    }


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("Usage: extract_pptx.py <input.pptx>\n")
        sys.exit(1)

    pptx_path = sys.argv[1]
    if not os.path.exists(pptx_path):
        sys.stderr.write(f"File not found: {pptx_path}\n")
        sys.exit(1)

    result = extract_pptx(pptx_path)
    print(json.dumps(result, ensure_ascii=False))
    sys.exit(0)


if __name__ == "__main__":
    main()
